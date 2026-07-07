# -*- coding: utf-8 -*-
"""일회성: OCR된 책 목차(위키 _index)에서 소단원 추출 → 진도 CSV + xlsx(5시트, 변시까지 회독 누계).
포맷: 소단원=추적단위. 없는 과목(상법·행정법·형소·선택·가족)은 생략. (2026-06-18)"""
import re, csv, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))  # noqa: E402
from _vault import VAULT_ROOT, vp  # noqa: E402

BASE = vp("sync", "위키", "원문")
OUTDIR = vp("6.진도관리", "data")
os.makedirs(OUTDIR, exist_ok=True)
# 과목 → 소단원 기준서(OCR된 1차 목차서)
BOOKS = [("민법", "쟁점노트_재산법"), ("민사소송법", "논점민소"),
         ("형법", "반반형법"), ("헌법", "헌법핵심정리300")]
NOISE = re.compile(r"판례색인|저자|약력|저서|신청서|색인|차례|개정|머리말|서문")
LINE = re.compile(r"^\d+\.\s*\[\[[^\|]+\|(.+?)\]\]\s*\(p\.(\d+)-(\d+)\)")
DAE = re.compile(r"^제\s*\d+\s*[편장]")

rows = []  # [과목, 大단원, 소단원, 책, 시작, 끝]
for subj, book in BOOKS:
    cur = ""
    path = os.path.join(BASE, book, "_index.md")
    for ln in open(path, encoding="utf-8").read().splitlines():
        m = LINE.match(ln)
        if not m:
            continue
        name, s, e = m.group(1).strip(), int(m.group(2)), int(m.group(3))
        if NOISE.search(name):
            continue
        if DAE.match(name):
            cur = re.sub(r"\s+", " ", name)
        rows.append([subj, cur, name, book, s, e])

# ---- CSV (소단원 마스터) ----
csvp = os.path.join(OUTDIR, "진도_소단원_master.csv")
with open(csvp, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(["과목", "大단원", "소단원", "출처책", "시작쪽", "끝쪽"])
    w.writerows(rows)

# ---- xlsx (트래커) ----
wb = openpyxl.Workbook()
HF = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1D9E75")
THIN = Border(*[Side(style="thin", color="DDDDDD")] * 4)
CEN = Alignment(horizontal="center", vertical="center")


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.font = HF; cell.fill = HFILL; cell.alignment = CEN
    ws.freeze_panes = "A2"


# 1) 안내
ws0 = wb.active; ws0.title = "안내"
guide = [
    ["진도 트래커 (수동입력) — 2026-06-18"],
    [""],
    ["단위", "소단원(교재 목차 기준). 추적·회독은 여기서."],
    ["소스", "민법=쟁점노트_재산법 / 민소=논점민소 / 형법=반반형법 / 헌법=헌법핵심정리300 (OCR된 책 목차)"],
    ["생략", "상법·행정법·형사소송법·선택법·가족법 = OCR 없어 제외(추후 OCR시 추가)"],
    [""],
    ["사용법", "[소단원진도] 시트에서 각 소단원 회독1~7 칸에 회독 날짜/O 입력 → 누계 자동. 상태=예정/진행중/종료."],
    ["", "[과목요약]은 자동 집계(진행률·평균회독). [주간로그]에 날짜별 기록 → 변시까지 누계."],
    ["변시 누계", "회독1~7 칸이 회독 누계. 주간로그가 시간축 누계. 회독 더 필요시 열 추가."],
]
for r in guide:
    ws0.append(r)
ws0.column_dimensions["A"].width = 12
ws0.column_dimensions["B"].width = 100
ws0["A1"].font = Font(bold=True, size=14)

# 2) 소단원진도 (메인)
ws = wb.create_sheet("소단원진도")
hd = ["과목", "大단원", "소단원", "책", "쪽수", "회독1", "회독2", "회독3", "회독4",
      "회독5", "회독6", "회독7", "누계", "상태", "메모"]
ws.append(hd)
for subj, dae, so, book, s, e in rows:
    ws.append([subj, dae, so, book, f"{s}-{e}", "", "", "", "", "", "", "", None, "", ""])
for i in range(2, ws.max_row + 1):
    ws.cell(i, 13).value = f"=COUNTA(F{i}:L{i})"      # 누계
style_header(ws, len(hd))
widths = [11, 22, 46, 14, 9, 7, 7, 7, 7, 7, 7, 7, 6, 8, 20]
for i, wd in enumerate(widths, 1):
    ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = wd
ws.auto_filter.ref = f"A1:O{ws.max_row}"

# 3) 과목요약 (자동집계 + 차트)
ws2 = wb.create_sheet("과목요약")
ws2.append(["과목", "소단원수", "종료", "진행중", "평균회독", "진행률"])
subs = ["민법", "민사소송법", "형법", "헌법"]
for r, subj in enumerate(subs, start=2):
    ws2.append([
        subj,
        f"=COUNTIF('소단원진도'!A:A,A{r})",
        f"=COUNTIFS('소단원진도'!A:A,A{r},'소단원진도'!N:N,\"종료\")",
        f"=COUNTIFS('소단원진도'!A:A,A{r},'소단원진도'!N:N,\"진행중\")",
        f"=IFERROR(AVERAGEIF('소단원진도'!A:A,A{r},'소단원진도'!M:M),0)",
        f"=IFERROR(C{r}/B{r},0)",
    ])
    ws2.cell(r, 6).number_format = "0%"
style_header(ws2, 6)
for col, wd in zip("ABCDEF", [13, 9, 7, 8, 9, 9]):
    ws2.column_dimensions[col].width = wd
ch = BarChart(); ch.title = "과목별 진행률"; ch.type = "bar"; ch.height = 7; ch.width = 14
data = Reference(ws2, min_col=6, min_row=1, max_row=5)
cats = Reference(ws2, min_col=1, min_row=2, max_row=5)
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
ws2.add_chart(ch, "H2")

# 4) 주간로그 (시간축 누계)
ws3 = wb.create_sheet("주간로그")
ws3.append(["날짜", "과목", "단원범위", "회독차", "분량(쪽/문항)", "비고"])
style_header(ws3, 6)
for col, wd in zip("ABCDEF", [12, 12, 40, 8, 14, 30]):
    ws3.column_dimensions[col].width = wd

# 5) 로드맵
ws4 = wb.create_sheet("로드맵")
for r in [["구분", "목표"], ["전체(변시까지)", ""], ["이번 방학", ""], ["다음 학기", ""]]:
    ws4.append(r)
style_header(ws4, 2)
ws4.column_dimensions["A"].width = 16; ws4.column_dimensions["B"].width = 80

xlsxp = os.path.join(OUTDIR, "진도_tracker.xlsx")
wb.save(xlsxp)
print(f"소단원 {len(rows)}개")
from collections import Counter
for k, v in Counter(r[0] for r in rows).items():
    print(f"  {k}: {v}")
print(f"CSV  → {csvp}")
print(f"xlsx → {xlsxp}")
